"""Software fixtures only: temporary unit arithmetic/boundary inputs are not experiments."""
import io
import json
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from jsonschema import validate as validate_json_schema
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest

from data_manager import DataManager, DataValidationError, ImportFormatError, FIELDS, UNITS, json_schema
from data_manager.importers import RawTable, read_table
from data_manager.validation import validate

ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / "data/templates/demo_schema_only.csv"


def stub(identifier="SCHEMA-TEST-001", **fields):
    return {"alloy": "6a01", "experiment_id": identifier, "source": "software_test_only",
            "record_kind": "schema_test", **fields}


def payload(*records, units=None):
    data = list(records) if units is None else {"records": list(records), "units": units}
    return json.dumps(data, ensure_ascii=False).encode("utf-8")


def preview(tmp_path, *records, **options):
    return DataManager(tmp_path).preview_bytes(payload(*records), "test.json", dataset="demo", **options)


def test_demo_import_no_measurements_and_no_research_pollution(tmp_path):
    manager = DataManager(tmp_path)
    assert manager.summary()["record_count"] == 0
    assert not manager.db_path.exists()
    before = manager.preview_bytes(DEMO.read_bytes(), DEMO.name, dataset="demo")
    assert before.ok and before.input_count == 1
    assert not manager.db_path.exists()
    metadata = manager.import_file(DEMO, dataset="demo")
    assert metadata["dataset_version"] == "v000001"
    assert metadata["record_count"] == metadata["added_count"] == 1
    assert metadata["updated_at"] and metadata["schema_version"] == "1.0.0"
    assert metadata["file_sha256"]
    record = manager.records("demo")[0]
    assert all(record[field] is None for field in UNITS)
    assert record["crack"] is None and record["process"] is None
    normalized = {key: value for key, value in record.items() if key not in ("input_row", "imported_version")}
    validate_json_schema(normalized, json_schema())
    assert manager.original_file("demo", 1) == DEMO.read_bytes()
    assert manager.summary()["record_count"] == 0
    assert manager.summary("demo")["overall_completeness"] == 17.65
    assert manager.summary("demo")["process_distribution"] == {"未提供": 1}
    assert manager.version_detail("demo", 1)["report"]["warning_count"] == 14
    with pytest.raises(DataValidationError):
        manager.import_file(DEMO, dataset="research")
    with pytest.raises(DataValidationError) as error:
        manager.import_file(DEMO, dataset="demo")
    assert error.value.report.issue_counts["duplicate"] == 1
    assert len(manager.versions("demo")) == 1


@pytest.mark.parametrize("field,value,unit,expected", [
    ("thickness", 1, "cm", 10), ("current", 1, "kA", 1000), ("voltage", 1, "kV", 1000),
    ("speed", 60, "mm/min", 1), ("speed", .06, "m/min", 1),
    ("heat_input", 1, "J/mm", .001), ("heat_input", 1, "kJ/cm", .1),
    ("tensile_strength", 1, "GPa", 1000), ("hardness", 1, "HV", 1),
    ("elongation", .1, "fraction", 10), ("porosity", 0, "%", 0), ("undercut", 0, "mm", 0),
])
def test_unit_conversion(tmp_path, field, value, unit, expected):
    report = preview(tmp_path, stub(**{field: {"value": value, "unit": unit}}))
    assert report.ok
    assert report.records[0][field] == pytest.approx(expected)


@pytest.mark.parametrize("fields,code", [
    ({"thickness": "1"}, "missing_unit"), ({"thickness": "1 kg"}, "invalid_unit"),
    ({"hardness": "1 HRC"}, "invalid_unit"), ({"speed": "-1 mm/s"}, "out_of_range"),
    ({"thickness": "0 mm"}, "out_of_range"), ({"porosity": "101 %"}, "out_of_range"),
    ({"current": float("inf"), "current_unit": "A"}, "invalid_number"),
    ({"current": True, "current_unit": "A"}, "invalid_number"),
    ({"current": "=1+1", "current_unit": "A"}, "invalid_number"),
    ({"source": "=1+1"}, "formula"), ({"crack": "maybe"}, "invalid_category"),
    ({"experiment_id": ""}, "missing_required"), ({"experiment_id": 1}, "invalid_type"),
    ({"unexpected": 1}, "unknown_field"),
    ({"thickness[mm]": "1 cm"}, "unit_conflict"),
    ({"thickness": "1 mm", "thickness[mm]": 1}, "duplicate_column"),
    ({"current": "1e309 A"}, "invalid_number"),
    ({"thickness": "1e308 m"}, "out_of_range"),
    ({"current": "1 MA"}, "invalid_unit"), ({"tensile_strength": "1 mPa"}, "invalid_unit"),
])
def test_invalid_data_never_creates_version(tmp_path, fields, code):
    manager = DataManager(tmp_path)
    with pytest.raises(DataValidationError) as error:
        manager.import_bytes(payload(stub(**fields)), "check.json", dataset="demo")
    assert error.value.report.issue_counts[code] >= 1
    assert manager.versions("demo") == []
    assert manager.summary("demo")["record_count"] == 0


def test_missing_unit_only_accepted_by_explicit_declaration(tmp_path):
    report = preview(tmp_path, stub(thickness=1), assume_canonical_units=True)
    assert report.ok and report.records[0]["thickness"] == 1
    assert report.issue_counts["assumed_unit"] == 1


def test_unit_column_and_global_declaration(tmp_path):
    manager = DataManager(tmp_path)
    csv = b'alloy,experiment_id,source,record_kind,speed,speed_unit\n6A01,TEST,software_test_only,schema_test,60,mm/min\n'
    report = manager.preview_bytes(csv, "input.csv", dataset="demo")
    assert report.ok and report.records[0]["speed"] == 1
    report = manager.preview_bytes(payload(stub(speed=60), units={"speed": "mm/min"}), "input.json", dataset="demo")
    assert report.ok and report.records[0]["speed"] == 1
    report = manager.preview_bytes(payload(stub(speed="60 mm/s"), units={"speed": "mm/min"}), "input.json", dataset="demo")
    assert not report.ok and report.issue_counts["unit_conflict"]


def test_xlsx_sheet_selection_formula_and_text_ids(tmp_path):
    book = Workbook()
    book.active.title = "instructions"
    book.active.append(["说明，不是实验表"])
    tab = book.create_sheet("records")
    tab.append(["alloy", "experiment_id", "source", "record_kind", "thickness[mm]"])
    tab.append(["6A01", "0001", "software_test_only", "schema_test", None])
    file = io.BytesIO()
    book.save(file)
    manager = DataManager(tmp_path)
    metadata = manager.import_bytes(file.getvalue(), "test.xlsx", dataset="demo", sheet="records")
    assert metadata["sheet"] == "records"
    assert manager.records("demo")[0]["experiment_id"] == "0001"
    with pytest.raises(ImportFormatError):
        manager.preview_bytes(file.getvalue(), "test.xlsx", dataset="demo", sheet="absent")
    tab.cell(2, 5, "=1+1")
    file = io.BytesIO()
    book.save(file)
    report = manager.preview_bytes(file.getvalue(), "test.xlsx", dataset="demo", sheet="records")
    assert report.issue_counts["invalid_number"]


@pytest.mark.parametrize("content,filename", [
    (b'', 'empty.csv'), (b'alloy,alloy\n6A01,6A01', 'duplicate.csv'),
    (b'alloy\n6A01,extra', 'extra.csv'), (b'{', 'broken.json'),
    (b'[{"alloy":"6A01","alloy":"6061"}]', 'duplicate.json'),
    (b'{"records": [], "unknown": 1}', 'unknown.json'),
    (b'[1]', 'invalid.json'), (b'bad', 'bad.xlsx'), (b'bad', 'bad.xls'), (b'bad', 'bad.txt'),
])
def test_malformed_files(content, filename):
    with pytest.raises(ImportFormatError):
        read_table(content, filename)


def test_empty_headers_and_gb18030(tmp_path):
    manager = DataManager(tmp_path)
    report = manager.preview_bytes(b'alloy,experiment_id,source\n', "empty.csv")
    assert not report.ok and report.issue_counts["empty_dataset"]
    csv = 'alloy,experiment_id,source,record_kind\n6A01,测试编号,软件测试,schema_test\n'.encode("gb18030")
    report = manager.preview_bytes(csv, "gb.csv", dataset="demo", encoding="gb18030")
    assert report.ok and report.records[0]["experiment_id"] == "测试编号"


def test_versions_batch_rollback_filters_pagination_and_zero_completeness(tmp_path):
    manager = DataManager(tmp_path)
    manager.import_bytes(payload(stub(porosity="0 %")), "first.json", dataset="demo")
    first = manager.version_detail("demo", 1)
    manager.import_bytes(payload(stub("SCHEMA-TEST-002", alloy="6061", process="TIG")), "second.json", dataset="demo")
    assert manager.version_detail("demo", 1) == first
    assert manager.summary("demo", version=1)["record_count"] == 1
    assert manager.summary("demo")["record_count"] == 2
    with pytest.raises(LookupError):
        manager.summary("demo", version=99)
    with pytest.raises(LookupError):
        manager.records("demo", version=99)
    assert manager.summary("demo", version=1)["completeness"]["porosity"] == 100
    assert manager.summary("demo")["completeness"]["porosity"] == 50
    assert manager.summary("demo")["alloy_distribution"] == {"6061": 1, "6A01": 1}
    assert manager.records("demo", offset=1, limit=1)[0]["experiment_id"] == "SCHEMA-TEST-002"
    assert manager.summary("demo", alloy="6061", process="TIG")["record_count"] == 1
    with pytest.raises(DataValidationError):
        manager.import_bytes(payload(stub("SCHEMA-TEST-003"), stub()), "mixed.json", dataset="demo")
    assert manager.summary("demo")["record_count"] == 2 and len(manager.versions("demo")) == 2
    with pytest.raises(DataValidationError):
        manager.import_bytes(payload(stub("NEW"), stub("NEW")), "duplicate.json", dataset="demo")
    assert manager.summary("demo")["record_count"] == 2


def test_identical_measurements_are_flagged_but_parallel_records_retained(tmp_path):
    report = preview(tmp_path, stub(current="1 A"), stub("SECOND", current="1 A"))
    assert report.ok and report.issue_counts["possible_duplicate"] == 1


def test_statistical_outliers_check_existing_data_without_mutating_values(tmp_path):
    manager = DataManager(tmp_path)
    # Arithmetic fixture for IQR, stored only in pytest's disposable directory.
    rows = [stub(f"SOFTWARE-IQR-{i}", current="1 A") for i in range(8)]
    manager.import_bytes(payload(*rows), "arithmetic.json", dataset="demo")
    report = manager.preview_bytes(payload(stub("SOFTWARE-IQR-BOUNDARY", current="100 A")), "arithmetic2.json", dataset="demo")
    assert report.ok and report.issue_counts["statistical_outlier"] == 1
    assert report.records[0]["current"] == 100
    assert manager.summary("demo")["record_count"] == 8


def test_concurrent_duplicate_import_is_atomic(tmp_path):
    manager = DataManager(tmp_path)
    manager.import_bytes(payload(stub()), "initial.json", dataset="demo")
    def run():
        try:
            return manager.import_bytes(payload(stub("CONCURRENT")), "same.json", dataset="demo")["version"]
        except DataValidationError:
            return "duplicate"
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: run(), range(2)))
    assert sorted(map(str, results)) == ["2", "duplicate"]
    assert manager.summary("demo")["record_count"] == 2


def test_issue_cap_retains_complete_error_counts():
    table = RawTable([{"row": i, "values": {}} for i in range(1000)], {})
    report = validate(table)
    assert not report.ok
    assert report.summary()["error_count"] == 3000
    assert len(report.issues) == 2000 and report.summary()["issues_truncated"]


def test_ten_thousand_placeholder_records_and_pagination(tmp_path):
    manager = DataManager(tmp_path)
    # Only unique routing IDs and labels, with no welding or performance data.
    content = payload(*(stub(f"BATCH-SCHEMA-{i:05d}") for i in range(10000)))
    metadata = manager.import_bytes(content, "batch_schema_only.json", dataset="demo")
    assert metadata["record_count"] == 10000
    assert manager.summary("demo")["record_count"] == 10000
    assert manager.summary("demo")["completeness"]["current"] == 0
    assert len(manager.records("demo", limit=100, offset=9900)) == 100
    assert manager.records("demo", limit=100, offset=10000) == []
    assert manager.summary("research")["record_count"] == 0


def test_checked_in_json_schema_matches_contract():
    assert json.loads((ROOT / "data_manager/experiment.schema.json").read_text(encoding="utf-8")) == json_schema()


def test_data_page_empty_demo_and_historical_views(monkeypatch, tmp_path):
    manager = DataManager(tmp_path)
    monkeypatch.setattr("data_manager.DataManager", lambda: manager)
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=15).run()
    app.switch_page("pages/3_科研数据管理.py").run()
    assert not app.exception
    assert app.metric[0].value == "0" and app.metric[2].value == "尚无版本"
    manager.import_file(DEMO, dataset="demo")
    app.radio[0].set_value("demo").run()
    assert not app.exception
    assert app.metric[0].value == "1" and app.metric[1].value == "17.65%"
    manager.import_bytes(payload(stub("PAGE-SECOND")), "second.json", dataset="demo")
    app.run()
    next(s for s in app.selectbox if s.label == "查看数据版本").select(1).run()
    assert not app.exception and app.metric[0].value == "1"
    app.radio[0].set_value("research").run()
    assert not app.exception and app.metric[0].value == "0"


def test_data_page_upload_commit_and_duplicate_block(monkeypatch, tmp_path):
    manager = DataManager(tmp_path)
    monkeypatch.setattr("data_manager.DataManager", lambda: manager)
    file = io.BytesIO(DEMO.read_bytes())
    file.name = DEMO.name
    monkeypatch.setattr("streamlit.file_uploader", lambda *args, **kwargs: file)
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=15).run()
    app.switch_page("pages/3_科研数据管理.py").run()
    assert not app.exception
    assert next(b for b in app.button if b.label == "导入并创建版本").disabled
    app.radio[0].set_value("demo").run()
    button = next(b for b in app.button if b.label == "导入并创建版本")
    assert not button.disabled
    button.click().run()
    assert not app.exception and app.success
    assert manager.summary("demo")["record_count"] == 1
    app.run()
    assert not app.exception and next(b for b in app.button if b.label == "导入并创建版本").disabled
