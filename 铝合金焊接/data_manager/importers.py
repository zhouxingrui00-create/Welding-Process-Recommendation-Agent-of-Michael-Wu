"""Read raw tables without guessing measurements, identifiers, or units."""
from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass
from pathlib import Path


class ImportFormatError(ValueError):
    pass


@dataclass
class RawTable:
    rows: list[dict]
    units: dict
    sheet: str | None = None


def _table(rows) -> list[dict]:
    iterator = iter(rows)
    headers = next(iterator, None)
    if headers is None:
        raise ImportFormatError("文件为空，缺少表头。")
    headers = [str(h).strip() if h is not None else "" for h in headers]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or any(not h for h in headers) or len(set(headers)) != len(headers):
        raise ImportFormatError("表头为空或重复，请使用唯一字段名。")
    result = []
    for row_number, values in enumerate(iterator, 2):
        values = list(values)
        if not any(v is not None and str(v).strip() for v in values):
            continue
        if any(v is not None and str(v).strip() for v in values[len(headers):]):
            raise ImportFormatError(f"第 {row_number} 行包含无表头的数据。")
        values += [None] * max(0, len(headers) - len(values))
        result.append({"row": row_number, "values": dict(zip(headers, values))})
    return result


def _unique_pairs(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ImportFormatError(f"JSON 重复键：{key}")
        result[key] = value
    return result


def read_table(content: bytes, filename: str, *, sheet: str | None = None,
               encoding: str = "utf-8-sig") -> RawTable:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".csv":
            return RawTable(_table(csv.reader(io.StringIO(content.decode(encoding)), strict=True)), {})
        if suffix == ".json":
            payload = json.loads(content.decode(encoding), object_pairs_hook=_unique_pairs)
            units = {}
            if isinstance(payload, dict):
                extra = set(payload) - {"records", "units"}
                if extra:
                    raise ImportFormatError(f"JSON 顶层未知字段：{sorted(extra)}")
                units = payload.get("units", {})
                payload = payload.get("records")
            if not isinstance(payload, list) or any(not isinstance(r, dict) for r in payload):
                raise ImportFormatError('JSON 应为记录数组或 {"records": [...], "units": {...}}。')
            if not isinstance(units, dict) or any(not isinstance(v, str) for v in units.values()):
                raise ImportFormatError("units 应为字段到单位字符串的映射。")
            return RawTable([{"row": i, "values": r} for i, r in enumerate(payload, 1)], units)
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                chosen = sheet or book.sheetnames[0]
                if chosen not in book.sheetnames:
                    raise ImportFormatError(f"不存在工作表：{chosen}")
                return RawTable(_table(book[chosen].iter_rows(values_only=True)), {}, chosen)
            finally:
                book.close()
        if suffix == ".xls":
            import xlrd
            book = xlrd.open_workbook(file_contents=content, on_demand=True)
            try:
                tab = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
                def rows():
                    for i in range(tab.nrows):
                        values = []
                        for cell in tab.row(i):
                            if cell.ctype in (xlrd.XL_CELL_ERROR, xlrd.XL_CELL_DATE):
                                raise ImportFormatError(f"第 {i+1} 行包含错误值或日期单元格，请使用原始文本/数值。")
                            values.append(bool(cell.value) if cell.ctype == xlrd.XL_CELL_BOOLEAN else cell.value)
                        yield values
                return RawTable(_table(rows()), {}, tab.name)
            finally:
                book.release_resources()
        raise ImportFormatError("仅支持 CSV、Excel (.xlsx/.xls) 和 JSON。")
    except ImportFormatError:
        raise
    except ImportError as exc:
        raise ImportFormatError("Excel 依赖未安装，请运行 pip install -r requirements.txt。") from exc
    except Exception as exc:
        raise ImportFormatError(f"读取 {filename} 失败：{exc}") from exc
