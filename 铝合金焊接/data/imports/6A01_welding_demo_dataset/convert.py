"""Read the supplied demo XLSX into isolated, review-only material JSON.

Uses openpyxl for reading only. No model training, live import or workbook edit.
"""
from copy import deepcopy
from hashlib import sha256
import json
import math
from pathlib import Path
import re
import shutil
from urllib.parse import urlparse

from openpyxl import load_workbook


OUT = Path(__file__).resolve().parent
PROJECT = OUT.parents[2]
INPUT = Path(r"C:\Users\22846\Desktop\焊接数据\6A01_welding_demo_dataset.xlsx")
SHEET = "6A01_demo_dataset"
HEADERS = ["ID", "Alloy", "Temper", "Process", "Joint / Other material", "Thickness_6A01_mm",
           "Filler", "Filler_diameter_mm", "Laser_power_kW", "Arc_current_A", "Wire_feed_m_min",
           "Welding_speed_m_min", "Gap_mm", "Shielding_gas", "Tensile_strength_MPa",
           "Fatigue_strength_MPa", "Weld_depth_mm", "Other_outcome", "Data_type",
           "Source_DOI_or_URL", "Notes"]
FIELDS = ("id", "alloy", "temper", "chemical_composition", "thickness_mm", "welding_method",
          "welding_parameters", "microstructure", "mechanical_properties", "defects", "source")
NUMERIC_PARAMETERS = {
    "Filler_diameter_mm": ("wire_diameter", "mm"),
    "Laser_power_kW": ("laser_power", "kW"),
    "Arc_current_A": ("arc_current", "A"),
    "Wire_feed_m_min": ("wire_feed_speed", "m/min"),
    "Welding_speed_m_min": ("welding_speed", "m/min"),
    "Gap_mm": ("gap", "mm"),
}
TYPE_MAP = {
    "Paper: measured": "paper_measurement_summary",
    "Paper: measured optimum": "paper_reported_optimum",
    "Paper: experiment/dataset setting": "paper_dataset_setting",
    "Paper: measured/model validation": "paper_measurement_or_model_validation",
    "Patent: example": "patent_example",
}


def dump(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def source_key(url):
    parsed = urlparse(url)
    if parsed.netloc.lower() in ("doi.org", "dx.doi.org"):
        return parsed.path.lstrip("/").lower(), "journal_article"
    match = re.fullmatch(r"/patent/(CN\d+[A-Z]\d?)/zh", parsed.path)
    if parsed.netloc.lower() == "patents.google.com" and match:
        return match.group(1), "patent"
    raise ValueError(f"未识别的来源格式：{url}")


def convert():
    input_hash = sha256(INPUT.read_bytes()).hexdigest()
    wb = load_workbook(INPUT, data_only=False)
    ws = wb[SHEET]
    assert [cell.value for cell in ws[1]] == HEADERS, "列结构变化，请先确认字段映射。"
    assert ws.max_row == 8 and ws.max_column == 21
    assert not any(cell.data_type == "f" for row in ws for cell in row), "需要先确认公式值，不能直接作为实测值读取。"
    checks = {row["key"]: row for row in json.loads((OUT / "source_checks.json").read_text(encoding="utf-8"))["records"]}
    source_copy = OUT / "source" / INPUT.name
    cases, properties, literature = [], [], {}

    for row in range(2, ws.max_row + 1):
        raw = {header: ws.cell(row, col).value for col, header in enumerate(HEADERS, 1)}
        cells = {header: f"{SHEET}!{ws.cell(row, col).coordinate}" for col, header in enumerate(HEADERS, 1)}
        assert raw["Alloy"] == "6A01"
        assert isinstance(raw["ID"], int) and not isinstance(raw["ID"], bool)
        assert raw["Data_type"] in TYPE_MAP
        for header in (*NUMERIC_PARAMETERS, "Thickness_6A01_mm", "Tensile_strength_MPa", "Fatigue_strength_MPa", "Weld_depth_mm"):
            value = raw[header]
            assert value is None or (isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)), (row, header)

        key, publication_type = source_key(raw["Source_DOI_or_URL"])
        check = checks.get(key)
        assert check is not None
        lit_id = "6A01-demo-literature-" + sha256(key.encode()).hexdigest()[:12]
        case_id = f"6A01-demo-case-{raw['ID']:03d}"
        origin = {"file_path": source_copy.relative_to(PROJECT).as_posix(), "file_sha256": input_hash,
                  "sheet": SHEET, "row": row, "range": f"A{row}:U{row}"}
        src = {"title": check["title"], "authors": check["authors"], "year": check["year"],
               "url": raw["Source_DOI_or_URL"], "doi": key if publication_type == "journal_article" else None,
               "report_id": key if publication_type == "patent" else None, "page": None,
               "publication_type": publication_type, **origin,
               "metadata_source_url": check["checked_url"],
               "note": "数值直接来自指定 Excel，文献书目信息另经公开来源匹配；尚未完成逐字段全文审核。"}
        if publication_type == "patent":
            src["inventors"] = check.get("inventors")

        def quantity(header, unit):
            if raw[header] is None:
                return None
            return {"value": raw[header], "unit": unit, "source_cell": cells[header],
                    "verification_status": "imported_from_excel_pending_full_review"}

        record = dict.fromkeys(FIELDS)
        record.update(id=case_id, alloy=raw["Alloy"], temper=raw["Temper"],
                      thickness_mm=raw["Thickness_6A01_mm"], welding_method=raw["Process"],
                      source=src, literature_id=lit_id, record_kind=TYPE_MAP[raw["Data_type"]],
                      data_type=raw["Data_type"], example_data=publication_type == "patent",
                      joint={"description": raw["Joint / Other material"]},
                      raw_record=raw, source_cells=cells, notes=raw["Notes"],
                      other_outcome=raw["Other_outcome"],
                      verification={"status": "pending_full_review", "source_identity": check["status"],
                                    "all_numeric_fields_verified": False, "eligible_for_recommendation": False,
                                    "eligible_for_training": False})
        params = {name: quantity(header, unit) for header, (name, unit) in NUMERIC_PARAMETERS.items()}
        params["filler_material"] = (None if raw["Filler"] is None else
                                     {"value": raw["Filler"], "unit": None, "source_cell": cells["Filler"]})
        params["shielding_gas"] = (None if raw["Shielding_gas"] is None else
                                   {"value": raw["Shielding_gas"], "unit": None, "source_cell": cells["Shielding_gas"]})
        record["welding_parameters"] = params
        record["geometry_results"] = {"penetration_depth": quantity("Weld_depth_mm", "mm")}

        # Descriptive outcomes remain literal; only explicitly named measurements are parsed.
        outcome = raw["Other_outcome"] or ""
        if "minimum-hardness HAZ" in outcome:
            record["fracture_information"] = {"location": "minimum-hardness HAZ", "source_cell": cells["Other_outcome"]}
        if "fracture in weld" in outcome:
            record["fracture_information"] = {"location": "weld", "description": outcome, "source_cell": cells["Other_outcome"]}
        if "fracture along IMC layer" in outcome:
            record["fracture_information"] = {"location": "IMC layer", "source_cell": cells["Other_outcome"]}
        if "X-ray inspection" in outcome:
            record["defects"] = {"type": "porosity", "inspection_method": "X-ray",
                                 "observation": outcome, "porosity_fraction": None,
                                 "source_cell": cells["Other_outcome"]}
        if record["record_kind"] == "paper_dataset_setting":
            record["result_scope"] = "数据集实验设置，不能当作一条已测得性能或缺陷标签的独立样本。"
        if record["record_kind"] == "paper_measurement_or_model_validation":
            record["result_scope"] = "保留原表实测/模型验证混合标记；具体数值属于实测还是模拟需回查原图表。"

        mechanical = {"tensile_strength": quantity("Tensile_strength_MPa", "MPa"),
                      "fatigue_strength": quantity("Fatigue_strength_MPa", "MPa")}
        if any(value is not None for value in mechanical.values()):
            mechanical["test_conditions"] = {"standard": None, "stress_ratio": None, "cycles": None,
                                             "sample_count": None, "loading_mode": None}
            record["mechanical_properties"] = mechanical
            prop = {field: deepcopy(record[field]) for field in FIELDS}
            prop.update(id=f"6A01-demo-property-{raw['ID']:03d}", case_id=case_id,
                        literature_id=lit_id, record_kind="case_property_extract",
                        data_type=raw["Data_type"], verification=deepcopy(record["verification"]),
                        note="关联案例的同一组性能摘录，不是独立新增实验；按 case_id 去重。")
            prop["welding_parameters"] = None
            properties.append(prop)
            record["property_record_id"] = prop["id"]
        cases.append(record)

        if key not in literature:
            lit = dict.fromkeys(FIELDS)
            lit.update(id=lit_id, alloy="6A01", title=check["title"], source={
                name: deepcopy(value) for name, value in src.items()
                if name not in ("row", "range")},
                publication_type=publication_type, case_ids=[], spreadsheet_rows=[],
                data_types=[], welding_methods=[], summary=None,
                verification={"status": "source_identified_data_pending_review",
                              "source_identity": check["status"], "full_text_data_review_complete": False},
                metadata_provenance={"url": check["checked_url"], "note": check["note"]})
            literature[key] = lit
        lit = literature[key]
        lit["case_ids"].append(case_id)
        lit["spreadsheet_rows"].append(row)
        for name, value in (("data_types", raw["Data_type"]), ("welding_methods", raw["Process"])):
            if value not in lit[name]:
                lit[name].append(value)

    assert len(cases) == 7 and len({row["id"] for row in cases}) == 7
    assert len(properties) == 3 and len(literature) == 6
    for kind, records in (("welding_cases", cases), ("properties", properties), ("literature", list(literature.values()))):
        # Reuse the module's null templates, not any existing active records.
        template = json.loads((PROJECT / "data" / "materials" / "6A01" / f"{kind}.json").read_text(encoding="utf-8"))
        payload = {"schema_version": template.get("schema_version"), "record_template": template.get("record_template"),
                   "import_status": "review_only_not_loaded_by_app", "records": records}
        dump(OUT / "materials" / "6A01" / f"{kind}.json", payload)
    snapshot = {"source_file": INPUT.name, "file_sha256": input_hash, "worksheets": []}
    for sheet in wb:
        snapshot["worksheets"].append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                                       "cells": {cell.coordinate: cell.value for row in sheet for cell in row}})
    dump(OUT / "raw_records.json", snapshot)
    source_copy.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(INPUT, source_copy)
    assert sha256(INPUT.read_bytes()).hexdigest() == input_hash
    assert sha256(source_copy.read_bytes()).hexdigest() == input_hash
    print(json.dumps({"cases": len(cases), "properties": len(properties), "literature": len(literature),
                      "papers": 5, "patents": 1, "source_unchanged": True, "output": str(OUT)}, ensure_ascii=False))


if __name__ == "__main__":
    convert()
